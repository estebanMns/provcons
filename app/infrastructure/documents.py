"""Lectura de los archivos que suben los usuarios (catálogos y cotizaciones).

Gemini lee PDF e imágenes directamente, pero no hojas de cálculo: un .xlsx se
convierte aquí a texto (una línea por fila, celdas separadas por " | ") y se
manda como prompt de texto.
"""

import asyncio
from io import BytesIO

from openpyxl import load_workbook

from app.core.exceptions import ValidationError
from app.infrastructure.ai_client import AIClient
from app.core.logging import get_logger

logger = get_logger(__name__)

# Formatos que Gemini lee como archivo
_FILE_MIME_TYPES = {
    "pdf": "application/pdf",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "png": "image/png",
}
# Tope para no mandar hojas gigantes en un solo prompt
MAX_SPREADSHEET_ROWS = 2000


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def spreadsheet_to_text(file_bytes: bytes) -> str:
    """Todas las hojas del libro como texto. Las filas vacías se omiten."""
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as e:
        raise ValidationError("No se pudo leer el archivo de Excel; verifica que no esté dañado") from e

    lines: list[str] = []
    rows = 0
    for sheet in workbook.worksheets:
        sheet_lines = []
        for row in sheet.iter_rows(values_only=True):
            cells = [_cell_text(v) for v in row]
            if not any(cells):
                continue
            sheet_lines.append(" | ".join(cells).rstrip(" |"))
            rows += 1
            if rows >= MAX_SPREADSHEET_ROWS:
                break
        if sheet_lines:
            lines.append(f"## Hoja: {sheet.title}")
            lines.extend(sheet_lines)
        if rows >= MAX_SPREADSHEET_ROWS:
            logger.warning("Hoja de cálculo recortada a %s filas", MAX_SPREADSHEET_ROWS)
            break
    workbook.close()
    return "\n".join(lines)


async def extract_json_from_upload(
    ai_client: AIClient, filename: str | None, file_bytes: bytes, system_instruction: str,
) -> dict:
    """Pide a la IA el JSON estructurado de un archivo subido, eligiendo cómo
    enviárselo según su tipo."""
    if not file_bytes:
        raise ValidationError("El archivo está vacío")

    extension = (filename or "").rsplit(".", 1)[-1].lower() if "." in (filename or "") else ""

    if extension == "xlsx":
        # Leer el Excel es CPU puro (openpyxl): en un hilo, para que el event
        # loop siga atendiendo las demás peticiones mientras tanto
        text = await asyncio.to_thread(spreadsheet_to_text, file_bytes)
        if not text.strip():
            raise ValidationError("La hoja de cálculo no tiene datos")
        logger.info("Excel convertido a texto: %s (%s caracteres)", filename, len(text))
        return await ai_client.generate_json(
            prompt=f"Contenido de la hoja de cálculo «{filename}»:\n\n{text}",
            system_instruction=system_instruction,
        ) or {}

    if extension == "xls":
        raise ValidationError("El formato .xls (Excel 97-2003) no es compatible; guárdalo como .xlsx e intenta de nuevo")

    mime_type = _FILE_MIME_TYPES.get(extension)
    if mime_type is None:
        raise ValidationError("Formato no compatible. Sube un PDF, una imagen (JPG o PNG) o un Excel (.xlsx)")

    return await ai_client.generate_json_from_file(
        file_bytes=file_bytes, mime_type=mime_type, system_instruction=system_instruction,
    ) or {}
